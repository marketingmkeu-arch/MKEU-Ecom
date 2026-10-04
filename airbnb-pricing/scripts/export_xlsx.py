"""Optionaler Excel-Export der Preispläne (benötigt openpyxl: pip install openpyxl).

Aufruf (nach `dynpricing.cli run`): python3 scripts/export_xlsx.py
"""

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "output"
SHEETS = [
    ("Airbnb-Eingabe 2026", "airbnb_eingabe_2026*.csv"),
    ("Airbnb-Eingabe Sommer 2027", "airbnb_eingabe_2027*.csv"),
    ("Tagesplan 2026", "preisplan_2026*.csv"),
    ("Tagesplan Sommer 2027", "preisplan_2027*.csv"),
]
HEAD = PatternFill("solid", fgColor="DDE7F0")
EVENT = PatternFill("solid", fgColor="FFF2CC")


def add_csv_sheet(wb: Workbook, title: str, path: Path) -> None:
    ws = wb.create_sheet(title)
    with open(path, encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    for r_idx, row in enumerate(rows, start=1):
        ws.append([int(v) if v.isdigit() else v for v in row])
        if r_idx == 1:
            for c in ws[1]:
                c.font, c.fill = Font(bold=True), HEAD
        elif "Event" in rows[0] and row[rows[0].index("Event")]:
            for c in ws[r_idx]:
                c.fill = EVENT
    for i, col in enumerate(rows[0], start=1):
        width = max(len(str(r[i - 1])) for r in rows[: min(len(rows), 200)]) if rows else 10
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, pattern in SHEETS:
        matches = sorted(OUT.glob(pattern))
        if matches:
            add_csv_sheet(wb, title, matches[-1])
    rev = json.loads((OUT / "umsatzprognose.json").read_text(encoding="utf-8"))
    ws = wb.create_sheet("Umsatzprognose")
    ws.append(["Szenario", "Monat", "Nächte", "Gastumsatz EUR", "Auszahlung EUR"])
    for name, sc in rev["szenarien"].items():
        for month, m in sc["monate"].items():
            ws.append([name, month, m["naechte"], m["brutto"], m["auszahlung"]])
        ws.append([name, "SUMME", sc["naechte"], sc["brutto_gast"], sc["auszahlung"]])
    ws.append([])
    ws.append(["Vergleich 2027 (realistisch)", "", "Nächte", "Ø Gastpreis", "Auszahlung EUR"])
    for name, v in rev["vergleich_2027"].items():
        ws.append([name, "", v["naechte"], v["adr"], v["auszahlung"]])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    for col in "ABCDE":
        ws.column_dimensions[col].width = 26
    wb.save(OUT / "preisplan.xlsx")
    print(OUT / "preisplan.xlsx")


if __name__ == "__main__":
    main()
